# -*- coding: utf-8 -*-
import io, sys
try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

OUT = r"report_xlsx.txt"
out = io.StringIO()
def p(*a): out.write(" ".join("" if x is None else str(x) for x in a) + "\n")

def dump(path, max_rows=25):
    p("="*90)
    p("FILE:", path)
    p("="*90)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        p("ERROR opening:", e); p(""); return
    p("SHEETS:", wb.sheetnames)
    for si, ws in enumerate(wb.worksheets):
        p("-"*90)
        p("SHEET[%d] '%s'  max_row=%s max_col=%s" % (si, ws.title, ws.max_row, ws.max_column))
        p("-"*90)
        r = 0
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            # trim trailing empties
            while vals and vals[-1] == "": vals.pop()
            p("R%03d | %s" % (r, " | ".join(vals)))
            r += 1
            if r >= max_rows:
                p("... (%d+ rows total, truncated)" % ws.max_row)
                break
        if si == 0 and "Забрать" not in path:
            break  # for Ozon template only need sheet 1
    p("")

dump(r"Остатки_18.08.2026.xlsx", 25)
dump(r"Забрать.xlsx", 40)

open(OUT, "w", encoding="utf-8").write(out.getvalue())
print("OK written", OUT)
