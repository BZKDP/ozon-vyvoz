# -*- coding: utf-8 -*-
r"""Заполняет шаблон утилизации Ozon (Склад | SKU | Количество) по списку босса.

Использование: python -X utf8 make_util.py <API_KEY>
Данные: свежий /v2/analytics/stock_on_warehouses кабинета 2843272,
фильтр по 19 SKU из Downloads\Забрать.xlsx, количество = free_to_sell
(бронь не трогаем, чтобы не превысить остаток).
Результат: Утилизация_GrafLab.xlsx + сводка в консоль.
"""
import json, urllib.request, urllib.error, sys, time, collections
import openpyxl

CID = "2843272"
KEY = sys.argv[1].strip()
OUT = r"Утилизация_GrafLab.xlsx"

def api(path, body, tries=6):
    data = json.dumps(body).encode()
    for a in range(tries):
        req = urllib.request.Request("https://api-seller.ozon.ru"+path, data=data,
            headers={"Client-Id": CID, "Api-Key": KEY, "Content-Type": "application/json"}, method="POST")
        try:
            return json.loads(urllib.request.urlopen(req, timeout=30).read().decode())
        except urllib.error.HTTPError as e:
            if e.code == 429:
                time.sleep(6*(a+1)); continue
            print("HTTP", e.code, e.read().decode()[:200]); raise
        except Exception as e:
            print("retry", a, e); time.sleep(3)
    raise SystemExit("API не отвечает: "+path)

# список босса: SKU из колонки A
wb = openpyxl.load_workbook(r"Забрать.xlsx", data_only=True)
ws = wb["Лист1"]
boss_skus, boss_qty, boss_art = set(), {}, {}
for r in list(ws.iter_rows(values_only=True))[1:]:
    if not r[1] or (isinstance(r[1], str) and ("ИТОГО" in r[1] or "ещё" in r[1])): continue
    try: sku = int(float(r[0]))
    except: continue
    boss_skus.add(sku)
    boss_art[sku] = str(r[1]).strip().upper()
    try: boss_qty[sku] = int(float(r[2]))
    except: boss_qty[sku] = 0
print("Босс-лист: %d SKU, %d шт (срез 18.08)" % (len(boss_skus), sum(boss_qty.values())))

rows, offset = [], 0
while True:
    j = api("/v2/analytics/stock_on_warehouses", {"limit": 1000, "offset": offset, "warehouse_type": "ALL"})
    part = (j.get("result") or {}).get("rows") or []
    rows += part
    if len(part) < 1000: break
    offset += 1000
print("API строк:", len(rows))

# (склад, sku) -> free_to_sell
out_rows = []
tot = 0
per_sku = collections.Counter()
for r in rows:
    sku = int(r.get("sku") or 0)
    if sku not in boss_skus: continue
    free = int(r.get("free_to_sell_amount") or 0)
    if free <= 0: continue
    out_rows.append((r.get("warehouse_name"), sku, free))
    tot += free
    per_sku[sku] += free

out_rows.sort(key=lambda x: (x[0], -x[2]))
nwb = openpyxl.Workbook()
nws = nwb.active
nws.title = "Лист1"
nws.append(["Склад", "SKU", "Количество"])
for w, s, q in out_rows:
    nws.append([w, s, q])
nwb.save(OUT)

print("\nСтрок в файле: %d | Итого к утилизации: %d шт" % (len(out_rows), tot))
print("Файл: %s\n" % OUT)
print("%-12s %-11s %6s %6s" % ("Артикул", "SKU", "18.08", "сейчас"))
for sku in sorted(boss_skus, key=lambda s: -boss_qty.get(s, 0)):
    print("%-12s %-11d %6d %6d" % (boss_art.get(sku, "?"), sku, boss_qty.get(sku, 0), per_sku.get(sku, 0)))
