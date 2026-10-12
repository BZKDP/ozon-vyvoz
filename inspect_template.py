# -*- coding: utf-8 -*-
import openpyxl, os, io, collections, sys
import json

path = sys.argv[1] if len(sys.argv) > 1 else r"Остатки_18.08.2026 (1).xlsx"
OUTDIR = r"."
REPORT = os.path.join(OUTDIR, "inspect_report.txt")

def num(x):
    try: return int(round(float(x)))
    except: return 0

# boss articles
wb = openpyxl.load_workbook(r"Забрать.xlsx", data_only=True)
ws = wb["Лист1"]
boss = set()
for r in list(ws.iter_rows(values_only=True))[1:]:
    if not r[1] or (isinstance(r[1], str) and ("ИТОГО" in r[1] or "ещё" in r[1])): continue
    boss.add(str(r[1]).strip().upper())

t = openpyxl.load_workbook(path, data_only=True)
ws = t.active
allrows = list(ws.iter_rows(values_only=True))
data = allrows[5:]  # с R005

out = io.StringIO()
def p(*a): out.write(" ".join("" if x is None else str(x) for x in a) + "\n")

p("Файл:", os.path.basename(path))
p("Строк данных:", len([r for r in data if r and r[2]]))
by_art = collections.defaultdict(lambda: {"g":0,"h":0,"rows":0,"wh":set(),"types":set()})
tot_g = tot_h = 0
types_all = collections.Counter()
for r in data:
    if not r or not r[2]: continue
    a = str(r[2]).strip().upper()
    g = num(r[6]); h = num(r[7])
    by_art[a]["g"] += g; by_art[a]["h"] += h; by_art[a]["rows"] += 1
    by_art[a]["wh"].add(r[3])
    if len(r) > 11 and r[11]: by_art[a]["types"].add(str(r[11])); types_all[str(r[11])] += 1
    tot_g += g; tot_h += h

p("Всего кол-во: столбец G(промарк/не треб) =", tot_g, " столбец H(немарк) =", tot_h)
p("Типы остатков в файле:", dict(types_all))
p("")
p("Артикулов всего в файле:", len(by_art), " | из них боссовских:", len(set(by_art) & boss))
p("")
p("%-12s %6s %6s %5s  %-8s %s" % ("Артикул","G","H","строк","БОСС?","Типы"))
p("-"*90)
for a in sorted(by_art, key=lambda x:-(by_art[x]["g"]+by_art[x]["h"])):
    d = by_art[a]
    p("%-12s %6d %6d %5d  %-8s %s" % (a, d["g"], d["h"], d["rows"], "ДА" if a in boss else "нет", ", ".join(sorted(d["types"]))[:40]))
open(REPORT, "w", encoding="utf-8").write(out.getvalue())
print("OK arts=%d boss=%d G=%d H=%d rows=%d" % (len(by_art), len(set(by_art)&boss), tot_g, tot_h, len([r for r in data if r and r[2]])))
