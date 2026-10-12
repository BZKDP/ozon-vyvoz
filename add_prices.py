import os
# -*- coding: utf-8 -*-
# Добавляет колонку с ценой в отчёт о списанных товарах (seller-disposes-report.xlsx)
import json, urllib.request, urllib.error
import openpyxl
from openpyxl.styles import Font

SRC = r"seller-disposes-report.xlsx"
OUT = r"seller-disposes-report_с_ценами.xlsx"

STORES = [
    ("163276", os.environ.get("OZON_API_KEY", "")),
    ("2843272", os.environ.get("OZON_API_KEY", "")),
]
BASE = "https://api-seller.ozon.ru"

def call(cid, key, path, body):
    h = {"Client-Id": cid, "Api-Key": key, "Content-Type": "application/json"}
    req = urllib.request.Request(BASE + path, data=json.dumps(body).encode(), headers=h, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return json.loads(r.read().decode())
    except urllib.error.HTTPError as e:
        print("HTTP", e.code, path, e.read().decode()[:200])
        return None

wb = openpyxl.load_workbook(SRC)
ws = wb.active
HDR = 4
# найти колонки
cols = {}
for c in range(1, ws.max_column + 1):
    v = ws.cell(HDR, c).value
    if v:
        cols[str(v).strip()] = c
sku_col = cols["SKU"]
art_col = cols["Артикул"]

rows = []
for r in range(HDR + 1, ws.max_row + 1):
    sku = ws.cell(r, sku_col).value
    art = ws.cell(r, art_col).value
    if sku or art:
        rows.append((r, sku, art))
print("Строк с товарами:", len(rows))

offer_ids = sorted(set(str(a) for _, _, a in rows if a))
print("Уникальных артикулов:", len(offer_ids))

# цены: offer_id -> price; пробуем по магазинам по очереди
prices = {}       # offer_id -> (price, store)
sku_prices = {}   # sku -> (price, store)
for cid, key in STORES:
    missing = [o for o in offer_ids if o not in prices]
    if not missing:
        break
    cursor = ""
    # v5 отдаёт по фильтру offer_id до 1000 позиций
    resp = call(cid, key, "/v5/product/info/prices",
                {"filter": {"offer_id": missing, "visibility": "ALL"}, "limit": 1000, "cursor": ""})
    items = (resp or {}).get("items", [])
    for it in items:
        oid = it.get("offer_id")
        p = (it.get("price") or {})
        val = p.get("price") or p.get("marketing_price")
        if oid and val:
            prices[oid] = (float(val), cid)
        for f in ("sku",):
            s = it.get(f)
            if s and val:
                sku_prices[int(s)] = (float(val), cid)
    print(f"Магазин {cid}: найдено цен {len(items)}")

print("Итого цен по артикулу:", len(prices), "из", len(offer_ids))

# новая колонка
new_col = ws.max_column + 1
hc = ws.cell(HDR, new_col, "Цена, руб.")
src_font = ws.cell(HDR, sku_col).font
hc.font = Font(name=src_font.name, size=src_font.size, bold=src_font.bold)

not_found = []
for r, sku, art in rows:
    hit = prices.get(str(art)) or (sku_prices.get(int(sku)) if sku else None)
    if hit:
        ws.cell(r, new_col, hit[0])
        ws.cell(r, new_col).number_format = "#,##0.00"
    else:
        ws.cell(r, new_col, "не найден")
        not_found.append(str(art))

ws.column_dimensions[openpyxl.utils.get_column_letter(new_col)].width = 12
wb.save(OUT)
print("Сохранено:", OUT)
nf = sorted(set(not_found))
print("Не найдено цен:", len(nf))
for x in nf[:30]:
    print(" ", x)
