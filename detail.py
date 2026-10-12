# -*- coding: utf-8 -*-
import json, os, collections, csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

OUTDIR = r"."
XLSX = os.path.join(OUTDIR, "Детализация_вывоз_2843272.xlsx")
CSVF = os.path.join(OUTDIR, "Детализация_по_складам.csv")
MD   = os.path.join(OUTDIR, "detail_chat.md")

def num(x):
    try: return int(round(float(x)))
    except: return 0

def cluster_of(wh):
    w = (wh or "").upper().replace("FRESH_MAX_WMS_", "")
    G = [("Зарубеж (KZ/BY)",["АЛМАТ","АСТАНА","КАРАГАНД","МИНСК","ГОМЕЛ","БРЕСТ"]),
         ("Москва и МО",["ПУШКИНО","ДОМОДЕД","ПЕТРОВСК","ХОРУГВИНО","ЖУКОВСК","ГРИВНО","НОГИНСК","КОЛЕД","СОФЬИНО","ОРЕХОВ","ЭЛЕКТРО","ТВЕР","ТУЛА","РЯЗАН","КАЛУГ","ВЛАДИМИР","ЯРОСЛАВ","МОСК"]),
         ("Санкт-Петербург и СЗО",["СПБ","ПЕТЕРБУРГ","ПАРГОЛОВ","ШУШАР","БУГРЫ","КОЛПИНО"]),
         ("Юг",["РОСТОВ","БАТАЙСК","АКСАЙ","КРАСНОДАР","АДЫГЕ","ВОЛГОГРАД","НОВОРОССИЙСК","НЕВИННОМЫСС","МАХАЧКАЛА","СТАВРОПОЛ"]),
         ("Черноземье",["ВОРОНЕЖ","ЛИПЕЦК","БЕЛГОРОД","КУРСК"]),
         ("Приволжье",["КАЗАН","ЗЕЛЕНОДОЛЬ","САМАР","ТОЛЬЯТ","СМР","УФА","НИЖН","ПЕРМ","САРАТОВ","ОРЕНБУРГ","ИЖЕВСК","ПЕНЗ","УЛЬЯНОВ"]),
         ("Урал",["ЕКАТЕРИН","ЕКБ","БЕРЕЗОВ","ТЮМЕН","ЧЕЛЯБ","ЧЛБ"]),
         ("Сибирь",["НОВОСИБ","КОЛЬЦОВ","НСК","ОМСК","КРАСНОЯР","БАРНАУЛ","ТОМСК","КЕМЕРОВ"]),
         ("Дальний Восток",["ХАБАРОВ","ВЛАДИВОСТ","АРТЕМ"]),
         ("Калининград",["КАЛИНИНГРАД"]) ]
    for name, keys in G:
        if any(k in w for k in keys): return name
    return "Прочее"

# boss
wb = openpyxl.load_workbook(r"Забрать.xlsx", data_only=True)
ws = wb["Лист1"]
boss = {}
for r in list(ws.iter_rows(values_only=True))[1:]:
    if not r[1] or (isinstance(r[1], str) and ("ИТОГО" in r[1] or "ещё" in r[1])): continue
    a = str(r[1]).strip().upper()
    boss[a] = dict(sku=num(r[0]), name=str(r[9] or ""), qty=num(r[2]))
boss_arts = set(boss)

# stock rows
rows = json.load(open(os.path.join(OUTDIR, "stock2_raw.json"), encoding="utf-8"))
cl_art = collections.defaultdict(lambda: collections.Counter())      # cluster -> art -> qty
cl_total = collections.Counter()
wh_records = []                                                        # склад-level
name_by_art = {}
sku_by_art = {}
for r in rows:
    a = str(r.get("item_code")).strip().upper()
    if a not in boss_arts: continue
    q = int(r.get("free_to_sell_amount") or 0)
    if q <= 0: continue
    cl = cluster_of(r.get("warehouse_name"))
    cl_art[cl][a] += q; cl_total[cl] += q
    name_by_art[a] = r.get("item_name") or boss[a]["name"]
    sku_by_art[a] = r.get("sku")
    wh_records.append(dict(cluster=cl, warehouse=r.get("warehouse_name"), art=a,
                           name=r.get("item_name") or boss[a]["name"], sku=r.get("sku"), qty=q))

clusters = sorted(cl_total, key=lambda x: -cl_total[x])
grand = sum(cl_total.values())

# ---------- Excel ----------
wbk = openpyxl.Workbook()
HF = PatternFill("solid", fgColor="2F5496"); H = Font(bold=True, color="FFFFFF")
CF = PatternFill("solid", fgColor="D9E1F2"); CB = Font(bold=True, size=12)
b = Side(style="thin", color="BFBFBF"); BORD = Border(left=b,right=b,top=b,bottom=b)

# Sheet 1: grouped cluster -> товары
s = wbk.active; s.title = "Кластер-Товар"
s.append(["Кластер / Артикул", "Товар", "SKU", "Кол-во, шт"])
for c in range(1,5): s.cell(1,c).font=H; s.cell(1,c).fill=HF; s.cell(1,c).alignment=Alignment(horizontal="center")
for cl in clusters:
    s.append([f"▎{cl}", "", "", cl_total[cl]])
    hr = s.max_row
    for c in range(1,5): s.cell(hr,c).fill=CF; s.cell(hr,c).font=CB
    for a in sorted(cl_art[cl], key=lambda x:-cl_art[cl][x]):
        s.append([a, name_by_art.get(a, boss[a]["name"]), sku_by_art.get(a), cl_art[cl][a]])
        for c in range(1,5): s.cell(s.max_row,c).border=BORD
s.append(["ВСЕГО","","",grand]); s.cell(s.max_row,1).font=Font(bold=True); s.cell(s.max_row,4).font=Font(bold=True)
for col,w in zip("ABCD",[30,52,13,11]): s.column_dimensions[col].width=w
s.freeze_panes="A2"

# Sheet 2: склад-level flat
s2 = wbk.create_sheet("Склад-Товар")
s2.append(["Кластер","Склад","Артикул","Товар","SKU","Кол-во, шт"])
for c in range(1,7): s2.cell(1,c).font=H; s2.cell(1,c).fill=HF; s2.cell(1,c).alignment=Alignment(horizontal="center")
for rec in sorted(wh_records, key=lambda x:(-cl_total[x["cluster"]], x["cluster"], x["warehouse"], -x["qty"])):
    s2.append([rec["cluster"], rec["warehouse"], rec["art"], rec["name"], rec["sku"], rec["qty"]])
    for c in range(1,7): s2.cell(s2.max_row,c).border=BORD
for col,w in zip("ABCDEF",[22,26,12,50,13,11]): s2.column_dimensions[col].width=w
s2.freeze_panes="A2"
wbk.save(XLSX)

# ---------- CSV (склад-level) ----------
with open(CSVF, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f); w.writerow(["Кластер","Склад","Артикул","Товар","SKU","Кол-во"])
    for rec in sorted(wh_records, key=lambda x:(-cl_total[x["cluster"]], x["cluster"], x["warehouse"], -x["qty"])):
        w.writerow([rec["cluster"], rec["warehouse"], rec["art"], rec["name"], rec["sku"], rec["qty"]])

# ---------- Markdown for chat ----------
md = []
md.append(f"Всего к вывозу: **{grand} шт**, {len(boss)} SKU, {len(clusters)} кластеров\n")
for cl in clusters:
    md.append(f"### {cl} — {cl_total[cl]} шт")
    md.append("| Артикул | Товар | Шт |")
    md.append("|---|---|--:|")
    for a in sorted(cl_art[cl], key=lambda x:-cl_art[cl][x]):
        nm = (name_by_art.get(a, boss[a]["name"]) or "").replace("|","/")
        md.append(f"| {a} | {nm} | {cl_art[cl][a]} |")
    md.append("")
open(MD, "w", encoding="utf-8").write("\n".join(md))
print("OK clusters=%d grand=%d rows_wh=%d pairs=%d" % (
    len(clusters), grand, len(wh_records), sum(len(v) for v in cl_art.values())))
