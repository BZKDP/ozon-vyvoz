# -*- coding: utf-8 -*-
import json, os, collections, io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

OUTDIR = r"."
XLSX = os.path.join(OUTDIR, "Отчет_вывоз_2843272.xlsx")
LPU = 0.9  # оценка литров на единицу (флакон БАД); из данных Ozon ~0.92 л/шт

def num(x):
    try: return int(round(float(x)))
    except: return 0

# clusters
def cluster_of(wh):
    w = (wh or "").upper().replace("FRESH_MAX_WMS_", "")
    G = [("Зарубеж (KZ/BY)",["АЛМАТ","АСТАНА","КАРАГАНД","МИНСК","ГОМЕЛ","БРЕСТ"]),
         ("Москва и МО",["ПУШКИНО","ДОМОДЕД","ПЕТРОВСК","ХОРУГВИНО","ЖУКОВСК","ГРИВНО","НОГИНСК","КОЛЕД","СОФЬИНО","ОРЕХОВ","ЭЛЕКТРО","ТВЕР","ТУЛА","РЯЗАН","КАЛУГ","ВЛАДИМИР","ЯРОСЛАВ","МОСК"]),
         ("Санкт-Петербург и СЗО",["СПБ","ПЕТЕРБУРГ","ПАРГОЛОВ","ШУШАР","БУГРЫ","КОЛПИНО"]),
         ("Юг",["РОСТОВ","БАТАЙСК","АКСАЙ","КРАСНОДАР","АДЫГЕ","ВОЛГОГРАД","НОВОРОССИЙСК","НЕВИННОМЫСС","МАХАЧКАЛА","СТАВРОПОЛ"]),
         ("Черноземье",["ВОРОНЕЖ","ЛИПЕЦК","БЕЛГОРОД","КУРСК"]),
         ("Приволжье",["КАЗАН","ЗЕЛЕНОДОЛЬ","САМАР","ТОЛЬЯТ","СМР","УФА","НИЖН","ПЕРМ","САРАТОВ","ОРЕНБУРГ","ИЖЕВСК","ПЕНЗ","УЛЬЯНОВ"]),
         ("Урал",["ЕКАТЕРИН","ЕКБ","БЕРЕЗОВ","ТЮМЕН","ЧЕЛЯБ","ЧЛБ"]),
         ("Сибирь",["НОВОСИБ","КОЛЬЦОВ","НСК","ОМСК","КРАСНОЯР","БАРНАУЛ","ТОМСК","КЕМЕРОВ"]),
         ("Дальний Восток",["ХАБАРОВ","ВЛАДИВОСТ","АРТЕМ"]),
         ("Калининград",["КАЛИНИНГРАД"]),
         ("Зарубеж (KZ/BY)",["АЛМАТ","АСТАНА","КАРАГАНД","МИНСК"]) ]
    for name, keys in G:
        if any(k in w for k in keys): return name
    return "Прочее"

# boss data
wb = openpyxl.load_workbook(r"Забрать.xlsx", data_only=True)
ws = wb["Лист1"]
boss = {}
for r in list(ws.iter_rows(values_only=True))[1:]:
    if not r[1] or (isinstance(r[1], str) and ("ИТОГО" in r[1] or "ещё" in r[1])): continue
    a = str(r[1]).strip().upper()
    boss[a] = dict(sku=num(r[0]), qty=num(r[2]), skl=num(r[3]), storage=num(r[6]), name=str(r[9] or ""))
boss_arts = set(boss)

# acct2 stock
rows = json.load(open(os.path.join(OUTDIR, "stock2_raw.json"), encoding="utf-8"))
# aggregate only boss items
cl_units = collections.Counter()
cl_arts = collections.defaultdict(set)
cl_wh = collections.defaultdict(set)
mat = collections.defaultdict(lambda: collections.Counter())  # article -> cluster -> qty
for r in rows:
    a = str(r.get("item_code")).strip().upper()
    if a not in boss_arts: continue
    q = int(r.get("free_to_sell_amount") or 0)
    cl = cluster_of(r.get("warehouse_name"))
    cl_units[cl] += q; cl_arts[cl].add(a); cl_wh[cl].add(r.get("warehouse_name")); mat[a][cl] += q

total = sum(cl_units.values())
clusters = sorted(cl_units, key=lambda x: -cl_units[x])

tot_storage = sum(boss[a]["storage"] for a in boss)

# --- workbook ---
wbk = openpyxl.Workbook()
H = Font(bold=True, size=11, color="FFFFFF")
HF = PatternFill("solid", fgColor="2F5496")
TITLE = Font(bold=True, size=14)
b = Side(style="thin", color="BFBFBF"); BORD = Border(left=b, right=b, top=b, bottom=b)
money = '#,##0'

def style_header(ws, ncol, row=1):
    for c in range(1, ncol+1):
        cell = ws.cell(row, c); cell.font = H; cell.fill = HF
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = BORD

# Sheet 1 — Сводка
s1 = wbk.active; s1.title = "Сводка"
s1["A1"] = "План вывоза товаров со стока Ozon — аккаунт 2843272 (GrafLab)"; s1["A1"].font = TITLE
data = [
    ("", ""),
    ("Позиций к вывозу (SKU):", len(boss)),
    ("Всего единиц:", total),
    ("Складов задействовано:", sum(len(v) for v in cl_wh.values())),
    ("Кластеров:", len(clusters)),
    ("Стоимость хранения этих товаров:", f"{tot_storage:,} ₽/мес".replace(",", " ")),
    ("", ""),
    ("ГЛАВНОЕ:", "хранение жжёт 204 765 ₽/мес, товары не продаются (статус «снят»)."),
    ("База:", "Москва и МО. Стратегия: максимально дёшево."),
    ("Москва (1773 шт):", "САМОВЫВОЗ. 85% в 2 складах: Пушкино_1 (795) + Хоругвино (709)."),
    ("Остальные 8 кластеров РФ:", "КУРЬЕР на московский адрес (ПВЗ дороже: обработка до 45₽/шт)."),
    ("Дальний Восток + Сибирь (1286 шт):", "сверить «Прогноз»: дорогая магистраль → часть утилизировать."),
    ("Минск (36 шт):", "в РФ не вывозится — списать/оставить."),
    ("Окупаемость:", "разовый вывоз << 204 765 ₽/мес хранения → окупается сразу."),
    ("Точная цена:", "в «Прогнозе стоимости» внутри заявки (магистраль с 04.2026 — по маршруту)."),
]
row = 3
for k, v in data:
    s1.cell(row, 1, k).font = Font(bold=True); s1.cell(row, 2, v); row += 1
s1.column_dimensions["A"].width = 34; s1.column_dimensions["B"].width = 80

# Sheet 2 — По кластерам
s2 = wbk.create_sheet("По кластерам")
hdr = ["Приоритет","Кластер","Единиц","Доля","Артикулов","Складов","~Литры","Рекоменд. способ","Ориентир ₽ (без магистрали)"]
s2.append(hdr); style_header(s2, len(hdr))
def method_for(cl, units):
    if cl.startswith("Зарубеж"): return "НЕ вывозится (Минск) — списать"
    if cl == "Москва и МО": return "САМОВЫВОЗ (дёшево, база тут)"
    if cl.startswith("Дальний") or cl == "Сибирь": return "Курьер (сверить прогноз!)"
    return "Курьер (на моск. адрес)"
for i, cl in enumerate(clusters, 1):
    u = cl_units[cl]; lit = round(u*LPU)
    if cl == "Москва и МО" or cl.startswith("Зарубеж"):
        est = 0  # самовывоз/списание — без вызова и перевозки
    else:
        est = round(546 + lit*1.1)  # вызов курьера + перевозка (без магистрали и сборки)
    s2.append([i, cl, u, u/total, len(cl_arts[cl]), len(cl_wh[cl]), lit, method_for(cl,u), est])
s2.append(["", "ИТОГО", total, 1.0, len(boss), sum(len(v) for v in cl_wh.values()), round(total*LPU), "", ""])
for r in range(2, s2.max_row+1):
    s2.cell(r,3).number_format = money; s2.cell(r,4).number_format='0%'
    s2.cell(r,7).number_format = money; s2.cell(r,9).number_format = money
    for c in range(1,len(hdr)+1): s2.cell(r,c).border = BORD
for col,w in zip("ABCDEFGHI",[9,24,9,7,10,8,9,26,26]): s2.column_dimensions[col].width=w
s2.cell(s2.max_row,2).font = Font(bold=True)

# Sheet 3 — По товарам
s3 = wbk.create_sheet("По товарам")
hdr3 = ["#","Артикул","SKU Ozon","Товар","Кол-во, шт","Складов","Хранение ₽/мес","Кластеры (где лежит)"]
s3.append(hdr3); style_header(s3, len(hdr3))
for i, a in enumerate(sorted(boss, key=lambda x: -boss[x]["storage"]), 1):
    where = ", ".join("%s:%d" % (c, mat[a][c]) for c in sorted(mat[a], key=lambda x:-mat[a][x]))
    s3.append([i, a, boss[a]["sku"], boss[a]["name"], boss[a]["qty"], boss[a]["skl"], boss[a]["storage"], where])
s3.append(["","ИТОГО","","",sum(boss[a]["qty"] for a in boss),"",tot_storage,""])
for r in range(2, s3.max_row+1):
    s3.cell(r,5).number_format=money; s3.cell(r,7).number_format=money
    for c in range(1,len(hdr3)+1): s3.cell(r,c).border=BORD
for col,w in zip("ABCDEFGH",[4,12,13,34,11,8,15,60]): s3.column_dimensions[col].width=w
s3.cell(s3.max_row,2).font=Font(bold=True)

# Sheet 4 — Матрица товар×кластер
s4 = wbk.create_sheet("Матрица")
head = ["Артикул","Товар"] + clusters + ["ИТОГО"]
s4.append(head); style_header(s4, len(head))
for a in sorted(boss, key=lambda x:-boss[x]["qty"]):
    line = [a, boss[a]["name"]] + [mat[a].get(c,0) or "" for c in clusters] + [sum(mat[a].values())]
    s4.append(line)
foot = ["ИТОГО",""] + [cl_units[c] for c in clusters] + [total]
s4.append(foot)
for r in range(2, s4.max_row+1):
    for c in range(3, len(head)+1):
        s4.cell(r,c).number_format=money; s4.cell(r,c).border=BORD
s4.column_dimensions["A"].width=12; s4.column_dimensions["B"].width=30
for idx in range(len(clusters)+1):
    s4.column_dimensions[chr(ord("C")+idx)].width=13
s4.cell(s4.max_row,1).font=Font(bold=True)

wbk.save(XLSX)
print("OK clusters=%d total=%d storage=%d file=%s" % (len(clusters), total, tot_storage, XLSX))
for c in clusters:
    print("  %-22s %5d" % (c, cl_units[c]))
