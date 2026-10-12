# -*- coding: utf-8 -*-
import openpyxl, io, json, os

DL = r"."
OUTDIR = r"."
REPORT = os.path.join(OUTDIR, "plan_report.txt")
PLANX = os.path.join(OUTDIR, "План_вывоза.xlsx")

out = io.StringIO()
def p(*a): out.write(" ".join("" if x is None else str(x) for x in a) + "\n")

def num(x):
    try: return int(round(float(x)))
    except: return 0

# --- boss list (Лист1) ---
wb = openpyxl.load_workbook(os.path.join(DL, "Забрать.xlsx"), data_only=True)
ws = wb["Лист1"]
rows = list(ws.iter_rows(values_only=True))
items = []
for r in rows[1:]:
    if r[1] in (None, "", "ИТОГО") or (isinstance(r[1], str) and "ещё" in r[1]):
        continue
    items.append(dict(
        sku=num(r[0]), art=str(r[1]).strip(), qty=num(r[2]), skl=num(r[3]),
        sales=r[4], storage=num(r[6]), status=str(r[8] or ""), name=str(r[9] or "")
    ))

tot_qty = sum(i["qty"] for i in items)
tot_stor = sum(i["storage"] for i in items)
maxskl = max(i["skl"] for i in items)
p("="*80); p("СПИСОК БОССА (Забрать.xlsx, Лист1) — что вывозим"); p("="*80)
p("Позиций (SKU):", len(items))
p("Всего единиц (Шт):", tot_qty)
p("Суммарное хранение, руб/мес:", tot_stor)
p("Складов на позицию: от", min(i['skl'] for i in items), "до", maxskl)
p("Все статусы:", ", ".join(sorted(set(i["status"] for i in items))))
p("")

# --- Ozon Хабаровск template articles ---
wt = openpyxl.load_workbook(os.path.join(DL, "Остатки_18.08.2026.xlsx"), data_only=True)
wst = wt.active
tmpl_arts = set()
for r in list(wst.iter_rows(values_only=True))[5:]:
    if r[2]: tmpl_arts.add(str(r[2]).strip())
p("Артикулы в скачанном шаблоне Ozon (Хабаровск):", ", ".join(sorted(tmpl_arts)))
p("")

# --- sellable stock from API (stock_raw) ---
sell_arts = set()
try:
    raw = json.load(open(os.path.join(OUTDIR, "stock_raw.json"), encoding="utf-8"))
    rws = raw.get("result", {}).get("rows", []) if isinstance(raw, dict) else raw
    sell_arts = set(str(r.get("item_code")).strip() for r in rws if r.get("item_code"))
except Exception as e:
    p("(stock_raw.json недоступен:", e, ")")

boss_arts = set(i["art"] for i in items)
p("="*80); p("СВЕРКА"); p("="*80)
ov_tmpl = boss_arts & tmpl_arts
ov_sell = boss_arts & sell_arts
p("Артикулов босса, которые ЕСТЬ в скачанном шаблоне Хабаровска:", len(ov_tmpl), "->", sorted(ov_tmpl) or "НИ ОДНОГО")
p("Артикулов босса, которые ЕСТЬ в продаваемом остатке (API free_to_sell):", len(ov_sell), "->", sorted(ov_sell) or "НИ ОДНОГО")
p("")
p(">> Вывод: если совпадений ~0 — скачанный шаблон содержит НЕ ТЕ товары.")
p("   Снятые с продажи позиции не видны под фильтром 'Доступные к продаже'.")
p("")

# --- priority table ---
items_sorted = sorted(items, key=lambda x: -x["storage"])
p("="*80); p("ПРИОРИТЕТ ПО ХРАНЕНИЮ (топ = дороже всего лежит)"); p("="*80)
p("%-3s %-11s %-6s %-5s %10s  %s" % ("#","Артикул","Шт","Скл","Хран/мес","Товар"))
cum = 0
for n, i in enumerate(items_sorted, 1):
    cum += i["storage"]
    p("%-3d %-11s %-6d %-5d %10d  %s" % (n, i["art"], i["qty"], i["skl"], i["storage"], i["name"]))
p("-"*80)
p("ИТОГО: %d SKU, %d шт, %d руб/мес хранения" % (len(items), tot_qty, tot_stor))
open(REPORT, "w", encoding="utf-8").write(out.getvalue())

# --- build plan xlsx for the boss ---
pb = openpyxl.Workbook(); s = pb.active; s.title = "План вывоза"
hdr = ["№","Артикул","SKU Ozon","Товар","Кол-во (шт)","Складов","Хранение ₽/мес","Статус","Экономия/мес (нараст.)"]
s.append(hdr)
cum = 0
for n, i in enumerate(items_sorted, 1):
    cum += i["storage"]
    s.append([n, i["art"], i["sku"], i["name"], i["qty"], i["skl"], i["storage"], i["status"], cum])
s.append(["", "ИТОГО", "", "", tot_qty, "", tot_stor, "", ""])
# widths
for col, w in zip("ABCDEFGHI", [4,12,13,34,11,8,15,10,20]):
    s.column_dimensions[col].width = w
from openpyxl.styles import Font
for c in s[1]: c.font = Font(bold=True)
pb.save(PLANX)

print("OK items=%d units=%d storage=%d overlap_tmpl=%d overlap_sell=%d" % (
    len(items), tot_qty, tot_stor, len(ov_tmpl), len(ov_sell)))
